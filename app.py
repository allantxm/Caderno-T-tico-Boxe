"""
Caderno Tático de Boxe
Diário pessoal de técnicas, combos, defesas, erros de sparring e treinos.

Dados: ficam em memória durante a sessão e são sincronizados com um
arquivo JSON num repositório do GitHub.

Secrets (Streamlit):
  GITHUB_TOKEN, GITHUB_REPO         obrigatórios para usar a nuvem
  GITHUB_FILE, GITHUB_BRANCH        opcionais
  APP_PASSWORD                      opcional: pede senha para abrir o app

Fundo opcional: coloque boxing_background.png na mesma pasta do app.py.
"""

import base64
import calendar
import hmac
import inspect
import io
import json
import time
import unicodedata
import uuid
from datetime import date, datetime, timedelta
from html import escape
from pathlib import Path
from urllib.parse import quote

import pandas as pd
import requests
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    from zoneinfo import ZoneInfo

    FUSO = ZoneInfo("America/Sao_Paulo")
except Exception:  # tzdata ausente
    FUSO = None


# ============================================================
# CONFIGURAÇÃO DA PÁGINA
# ============================================================

st.set_page_config(
    page_title="Caderno Tático de Boxe",
    page_icon="🥊",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ============================================================
# CONSTANTES
# ============================================================

COLUNAS = [
    "id",
    "tipo",
    "titulo",
    "categoria",
    "descricao",
    "fonte",
    "link_video",
    "dificuldade",
    "favorito",
    "data",
]

# Cada tipo tem uma cor própria (a "fita" do cartão) e a chave usada no JSON.
TIPOS_INFO = {
    "Dica": {"css": "t-dica", "cor": "#5BA4E6", "chave": "anotacoes"},
    "Combo": {"css": "t-combo", "cor": "#F28C45", "chave": "combos"},
    "Defesa": {"css": "t-defesa", "cor": "#3DB5A6", "chave": "defesas"},
    "Erro no Sparring": {"css": "t-erro", "cor": "#E8C547", "chave": "erros_sparring"},
    "Correção do Professor": {"css": "t-correcao", "cor": "#A98BEA", "chave": "correcoes_professor"},
    "Treino": {"css": "t-treino", "cor": "#7CC47C", "chave": "treinos"},
}
TIPOS = list(TIPOS_INFO)

CATEGORIAS = [
    "Footwork / Movimentação",
    "Ataque & Combos",
    "Defesa & Esquivas",
    "Estratégia & Sparring",
    "Técnica",
    "Condicionamento",
    "Mentalidade",
    "Outro",
]

DIFICULDADES = ["Fácil", "Médio", "Difícil"]
PESO_DIFICULDADE = {"Fácil": 1, "Médio": 2, "Difícil": 3}
TIPOS_REVISAO = ["Erro no Sparring", "Correção do Professor"]

# (rótulo no menu, tipo filtrado ou None)
PAGINAS = [
    ("Visão geral", None),
    ("Caderno", None),
    ("Dicas", "Dica"),
    ("Combos", "Combo"),
    ("Defesas", "Defesa"),
    ("Erros no sparring", "Erro no Sparring"),
    ("Correções do professor", "Correção do Professor"),
    ("Treinos", "Treino"),
    ("Favoritos", None),
]

VALORES_PADRAO = {
    "id": "",
    "tipo": "Dica",
    "titulo": "",
    "categoria": "Outro",
    "descricao": "",
    "fonte": "",
    "link_video": "",
    "dificuldade": "Médio",
    "favorito": False,
    "data": "",
}

ORDENACOES = [
    "Mais recentes",
    "Mais antigos",
    "Título (A a Z)",
    "Mais difíceis primeiro",
    "Favoritos primeiro",
]

GITHUB_API = "https://api.github.com"
ITENS_POR_PAGINA = 10
ARQUIVO_FUNDO = "boxing_background.png"

MESES = [
    "janeiro", "fevereiro", "março", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
]
DIAS_SEMANA = [
    "segunda-feira", "terça-feira", "quarta-feira", "quinta-feira",
    "sexta-feira", "sábado", "domingo",
]

# Largura total de botões: o parâmetro mudou nas versões novas do Streamlit.
try:
    if "width" in inspect.signature(st.button).parameters:
        LARGURA_BTN = {"width": "stretch"}
    else:
        LARGURA_BTN = {"use_container_width": True}
except Exception:
    LARGURA_BTN = {"use_container_width": True}


# ============================================================
# UTILITÁRIOS
# ============================================================

def agora():
    """Data/hora no fuso de São Paulo (o servidor costuma rodar em UTC)."""
    return datetime.now(FUSO) if FUSO else datetime.now()


def hoje():
    return agora().date()


def novo_id():
    return uuid.uuid4().hex[:16]


def normalizar_texto(texto):
    """Minúsculas e sem acentos, para busca tolerante."""
    texto = unicodedata.normalize("NFKD", str(texto))
    return texto.encode("ascii", "ignore").decode("ascii").lower()


def normalizar_data(valor):
    """Converte vários formatos para AAAA-MM-DD (ou '' se inválido)."""
    texto = str(valor if valor is not None else "").strip()

    if not texto or texto.lower() in ("nan", "nat", "none"):
        return ""

    for formato in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(texto[:10], formato).strftime("%Y-%m-%d")
        except ValueError:
            continue

    try:
        return pd.to_datetime(texto).strftime("%Y-%m-%d")
    except Exception:
        return ""


def formatar_data(iso):
    if not iso:
        return ""
    try:
        return datetime.strptime(iso[:10], "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return str(iso)


def para_date(iso):
    try:
        return datetime.strptime(iso[:10], "%Y-%m-%d").date()
    except (ValueError, TypeError):
        return hoje()


def data_por_extenso(dia):
    return f"{DIAS_SEMANA[dia.weekday()].capitalize()}, {dia.day} de {MESES[dia.month - 1]}"


def normalizar_link(link):
    link = str(link or "").strip()
    if link and not link.lower().startswith(("http://", "https://")):
        link = "https://" + link
    return link


def plural(n, singular, plural_):
    return f"{n} {singular if n == 1 else plural_}"


def trecho(texto, limite=170):
    texto = " ".join(str(texto).split())
    return texto if len(texto) <= limite else texto[: limite - 1].rstrip() + "…"


def indice_de(lista, valor, padrao=0):
    return lista.index(valor) if valor in lista else padrao


# ============================================================
# DATAFRAME
# ============================================================

def criar_dataframe_vazio():
    return pd.DataFrame(columns=COLUNAS)


def converter_favorito(valor):
    if isinstance(valor, bool):
        return valor
    return str(valor).strip().lower() in ("true", "1", "sim", "yes", "favorito", "x")


def preparar_dataframe(df):
    """
    Garante colunas, tipos e valores válidos. Aceita dados antigos
    (sem 'descricao', com colunas a mais, etc.) e planilhas importadas.
    """
    if df is None:
        return criar_dataframe_vazio()

    df = df.copy()

    for coluna, padrao in VALORES_PADRAO.items():
        if coluna not in df.columns:
            df[coluna] = padrao

    df = df[COLUNAS].reset_index(drop=True)

    # Textos limpos
    for coluna in COLUNAS:
        if coluna == "favorito":
            continue
        df[coluna] = df[coluna].apply(lambda v: "" if pd.isna(v) else str(v).strip())

    # IDs únicos (gera para os ausentes e para os repetidos)
    vistos, ids = set(), []
    for valor in df["id"]:
        if not valor or valor in vistos:
            valor = novo_id()
        vistos.add(valor)
        ids.append(valor)
    df["id"] = ids

    # Valores de listas fechadas (ignora maiúsculas/minúsculas)
    def forcar(valor, opcoes, padrao):
        return {o.lower(): o for o in opcoes}.get(valor.lower(), padrao)

    df["tipo"] = df["tipo"].apply(lambda v: forcar(v, TIPOS, "Dica"))
    df["dificuldade"] = df["dificuldade"].apply(lambda v: forcar(v, DIFICULDADES, "Médio"))
    df["categoria"] = df["categoria"].apply(lambda v: v or "Outro")
    df["data"] = df["data"].apply(normalizar_data)
    df["favorito"] = df["favorito"].apply(converter_favorito).astype(bool)

    return df


def datas_de_treino(df):
    """Conjunto de datas (AAAA-MM-DD) que têm pelo menos um registro de Treino."""
    treinos = df[(df["tipo"] == "Treino") & (df["data"] != "")]
    return set(treinos["data"])


def sequencia_de_treinos(datas_iso, referencia):
    """
    Quantos dias seguidos o aluno treinou. Se hoje ainda não tem treino,
    a sequência continua valendo a partir de ontem.
    """
    dias = {para_date(d) for d in datas_iso}
    dia = referencia if referencia in dias else referencia - timedelta(days=1)
    total = 0
    while dia in dias:
        total += 1
        dia -= timedelta(days=1)
    return total


# ============================================================
# JSON <-> DATAFRAME
# ============================================================

def json_para_dataframe(dados):
    """Converte a estrutura JSON (listas por tipo) para um DataFrame único."""
    if not isinstance(dados, dict):
        return criar_dataframe_vazio()

    registros = []

    for tipo, info in TIPOS_INFO.items():
        lista = dados.get(info["chave"], [])

        if not isinstance(lista, list):
            continue

        for item in lista:
            if not isinstance(item, dict):
                continue
            registro = item.copy()
            if not registro.get("tipo"):
                registro["tipo"] = tipo
            registros.append(registro)

    if not registros:
        return criar_dataframe_vazio()

    return preparar_dataframe(pd.DataFrame(registros))


def dataframe_para_json(df):
    """Converte o DataFrame para o formato do arquivo JSON."""
    df = preparar_dataframe(df)
    resultado = {info["chave"]: [] for info in TIPOS_INFO.values()}

    for _, linha in df.iterrows():
        registro = {
            coluna: (bool(linha[coluna]) if coluna == "favorito" else str(linha[coluna]))
            for coluna in COLUNAS
        }
        resultado[TIPOS_INFO[registro["tipo"]]["chave"]].append(registro)

    return resultado


# ============================================================
# GITHUB
# ============================================================

def _segredo(nome, padrao=""):
    try:
        return st.secrets.get(nome, padrao)
    except Exception:  # sem secrets.toml
        return padrao


def obter_config_github():
    return (
        _segredo("GITHUB_TOKEN"),
        _segredo("GITHUB_REPO"),
        _segredo("GITHUB_FILE", "data/caderno_boxe.json"),
        _segredo("GITHUB_BRANCH"),
    )


def github_configurado():
    token, repo, _, _ = obter_config_github()
    return bool(token and repo)


def github_headers(token, aceitar="application/vnd.github+json"):
    return {
        "Authorization": f"Bearer {token}",
        "Accept": aceitar,
        "X-GitHub-Api-Version": "2022-11-28",
    }


def _url_arquivo(repo, arquivo):
    return f"{GITHUB_API}/repos/{repo}/contents/{quote(arquivo, safe='/')}"


def carregar_github():
    """
    Carrega o JSON do GitHub.
    Retorna (DataFrame, mensagem_de_erro, estado) com estado em:
    'ok', 'novo' (arquivo ainda não existe), 'erro' ou 'sem_config'.
    """
    token, repo, arquivo, branch = obter_config_github()

    if not token or not repo:
        return criar_dataframe_vazio(), "Nuvem não configurada (GITHUB_TOKEN e GITHUB_REPO).", "sem_config"

    try:
        resposta = requests.get(
            _url_arquivo(repo, arquivo),
            headers=github_headers(token, "application/vnd.github.raw+json"),
            params={"ref": branch} if branch else None,
            timeout=20,
        )

        if resposta.status_code == 404:
            return criar_dataframe_vazio(), "", "novo"

        if resposta.status_code in (401, 403):
            return (
                criar_dataframe_vazio(),
                "O GitHub recusou o acesso. Confira se o token é válido e tem permissão no repositório.",
                "erro",
            )

        resposta.raise_for_status()
        dados = json.loads(resposta.content.decode("utf-8"))
        return json_para_dataframe(dados), "", "ok"

    except Exception as erro:
        return criar_dataframe_vazio(), f"Não foi possível carregar da nuvem: {erro}", "erro"


def salvar_github(df):
    """Salva o DataFrame no JSON do GitHub. Retorna (sucesso, mensagem)."""
    token, repo, arquivo, branch = obter_config_github()

    if not token or not repo:
        return False, "Nuvem não configurada (GITHUB_TOKEN e GITHUB_REPO)."

    url = _url_arquivo(repo, arquivo)
    parametros = {"ref": branch} if branch else None

    try:
        # SHA atual do arquivo (necessário para atualizar)
        resposta_get = requests.get(
            url, headers=github_headers(token), params=parametros, timeout=20
        )

        sha = None
        if resposta_get.status_code == 200:
            sha = resposta_get.json().get("sha")
        elif resposta_get.status_code != 404:
            resposta_get.raise_for_status()

        texto_json = json.dumps(dataframe_para_json(df), ensure_ascii=False, indent=2)

        payload = {
            "message": f"Atualiza caderno ({agora():%d/%m/%Y %H:%M})",
            "content": base64.b64encode(texto_json.encode("utf-8")).decode("ascii"),
        }
        if sha:
            payload["sha"] = sha
        if branch:
            payload["branch"] = branch

        resposta_put = requests.put(
            url, headers=github_headers(token), json=payload, timeout=20
        )

        if resposta_put.status_code in (409, 422):
            return False, (
                "O arquivo mudou na nuvem enquanto você editava. "
                "Baixe um backup e recarregue da nuvem."
            )

        if resposta_put.status_code in (401, 403):
            return False, "O GitHub recusou a gravação. Confira a permissão de escrita do token."

        resposta_put.raise_for_status()
        return True, "Caderno salvo na nuvem."

    except Exception as erro:
        return False, f"Não foi possível salvar na nuvem: {erro}"


# ============================================================
# EXPORTAÇÃO
# ============================================================

def gerar_excel(df):
    """Gera o .xlsx em memória (cabeçalho estilizado, filtros e colunas ajustadas)."""
    export_df = preparar_dataframe(df)

    wb = Workbook()
    ws = wb.active
    ws.title = "Caderno Tático"
    ws.append(COLUNAS)

    for linha in export_df.itertuples(index=False, name=None):
        valores = dict(zip(COLUNAS, linha))
        valores["favorito"] = "Sim" if valores["favorito"] else "Não"
        valores["data"] = formatar_data(valores["data"])
        ws.append([valores[c] for c in COLUNAS])

    cabecalho_fill = PatternFill("solid", fgColor="162434")
    borda = Border(bottom=Side(style="thin", color="2A4057"))

    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = cabecalho_fill
        celula.alignment = Alignment(horizontal="left", vertical="center")

    for linha in ws.iter_rows(min_row=2):
        for celula in linha:
            celula.alignment = Alignment(vertical="top", wrap_text=True)
            celula.border = borda

    larguras = {
        "id": 18, "tipo": 22, "titulo": 36, "categoria": 26, "descricao": 70,
        "fonte": 26, "link_video": 40, "dificuldade": 13, "favorito": 10, "data": 12,
    }
    for posicao, coluna in enumerate(COLUNAS, start=1):
        ws.column_dimensions[get_column_letter(posicao)].width = larguras[coluna]

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def gerar_backup_json(df):
    return json.dumps(dataframe_para_json(df), ensure_ascii=False, indent=2).encode("utf-8")


# ============================================================
# ACESSO COM SENHA (opcional)
# ============================================================

def exigir_senha():
    """
    Se APP_PASSWORD estiver nos Secrets, nada é carregado nem exibido
    antes de a senha ser digitada.
    """
    senha = str(_segredo("APP_PASSWORD", ""))

    if not senha or st.session_state.get("autenticado"):
        return

    st.markdown(
        '<h1 class="page-title">Caderno tático</h1>'
        '<div class="page-sub">Digite a senha para abrir seu caderno.</div>',
        unsafe_allow_html=True,
    )

    with st.form("form_login"):
        tentativa = st.text_input("Senha", type="password")
        entrar = st.form_submit_button("Entrar", type="primary")

    if entrar:
        if hmac.compare_digest(tentativa.encode("utf-8"), senha.encode("utf-8")):
            st.session_state.autenticado = True
            st.rerun()

        # Pausa crescente a cada erro, para dificultar tentativas em sequência
        erros = st.session_state.get("tentativas_senha", 0) + 1
        st.session_state.tentativas_senha = erros
        time.sleep(min(erros, 5))
        st.error("Senha incorreta.")

    st.stop()


# ============================================================
# ESTADO E MUTAÇÕES
# ============================================================

def iniciar_estado():
    if "df" not in st.session_state:
        with st.spinner("Carregando seu caderno..."):
            df_inicial, erro, estado = carregar_github()

        st.session_state.df = preparar_dataframe(df_inicial)
        st.session_state.erro_conexao = erro
        st.session_state.estado_nuvem = estado
        st.session_state.alterado = False
        st.session_state.ultimo_salvamento = ""
        st.session_state.autosave = estado in ("ok", "novo")

    st.session_state.setdefault("pagina", "Visão geral")
    st.session_state.setdefault("mes_calendario", hoje().replace(day=1))
    st.session_state.setdefault("data_treino_selecionada", "")
    st.session_state.setdefault("editando_id", None)
    st.session_state.setdefault("excluindo_id", None)
    st.session_state.setdefault("avisos", [])
    st.session_state.setdefault("confirmar_recarga", False)


def ir_para(pagina, data_treino=""):
    """Troca de página limpando qualquer edição ou filtro de data pendente."""
    st.session_state.pagina = pagina
    st.session_state.data_treino_selecionada = data_treino
    st.session_state.editando_id = None
    st.session_state.excluindo_id = None


def avisar(texto, icone="✅"):
    """Guarda uma mensagem para aparecer depois do próximo rerun."""
    st.session_state.avisos.append((icone, texto))


def pode_salvar_na_nuvem():
    return github_configurado() and st.session_state.estado_nuvem in ("ok", "novo")


def salvar_na_nuvem(silencioso=False):
    """Salva agora. Retorna True se deu certo. Com silencioso=True, só avisa em caso de erro."""
    if not github_configurado():
        avisar("Configure GITHUB_TOKEN e GITHUB_REPO nos Secrets para usar a nuvem.", "⚠️")
        return False

    if not pode_salvar_na_nuvem():
        avisar(
            "O caderno não foi carregado da nuvem. Recarregue antes de salvar "
            "para não sobrescrever seus dados.",
            "⚠️",
        )
        return False

    sucesso, mensagem = salvar_github(st.session_state.df)

    if sucesso:
        st.session_state.alterado = False
        st.session_state.estado_nuvem = "ok"
        st.session_state.ultimo_salvamento = agora().strftime("%H:%M")
        if not silencioso:
            avisar(mensagem)
    else:
        avisar(mensagem, "⚠️")

    return sucesso


def confirmar_mudanca(df_novo, mensagem, icone="✅"):
    """Aplica uma mudança no caderno e salva sozinho, se estiver ativado."""
    st.session_state.df = preparar_dataframe(df_novo)
    st.session_state.alterado = True
    avisar(mensagem, icone)

    if st.session_state.autosave and pode_salvar_na_nuvem():
        salvar_na_nuvem(silencioso=True)


def atualizar_registro(id_registro, campos, mensagem, icone="✅"):
    df = st.session_state.df.copy()
    posicoes = df.index[df["id"] == id_registro]

    if len(posicoes) == 0:
        return

    for coluna, valor in campos.items():
        df.at[posicoes[0], coluna] = valor

    confirmar_mudanca(df, mensagem, icone)


def excluir_registro(id_registro):
    df = st.session_state.df
    df = df[df["id"] != id_registro].reset_index(drop=True)
    confirmar_mudanca(df, "Registro excluído.", "🗑️")


def adicionar_registro(campos, mensagem="Registro adicionado."):
    novo = {"id": novo_id(), **campos}
    df = pd.concat([st.session_state.df, pd.DataFrame([novo])], ignore_index=True)
    confirmar_mudanca(df, mensagem)


def marcar_dia_como_treinado(iso):
    """Cria um treino simples para o dia; os detalhes podem ser editados depois."""
    adicionar_registro(
        {
            "tipo": "Treino",
            "titulo": "Treino",
            "categoria": "Condicionamento",
            "descricao": "Treino marcado pelo calendário. Edite para adicionar os detalhes.",
            "fonte": "",
            "link_video": "",
            "dificuldade": "Médio",
            "favorito": False,
            "data": iso,
        },
        mensagem="Dia marcado como treinado.",
    )


def mostrar_avisos():
    avisos, st.session_state.avisos = st.session_state.avisos, []
    for icone, texto in avisos:
        st.toast(texto, icon=icone)


# ============================================================
# ESTILO
# ============================================================

# Painéis com contorno neon: o marcador invisível dentro do container
# identifica qual deles estilizar (via :has).
PAINEIS_NEON = ("hm", "calendar-marker", "recentes-marker")
CALENDARIO = "div[data-testid='stVerticalBlockBorderWrapper']:has(.calendar-marker)"

CSS_BASE = """
@import url('https://fonts.googleapis.com/css2?family=Barlow:wght@400;500;600&family=Barlow+Condensed:wght@500;600;700&display=swap');

:root{
  --bg:#000; --surface:#080808; --raised:#141414; --line:#292929;
  --text:#F5F5F5; --muted:#A7A7A7;
  --accent:#E1262F; --accent-hi:#FF303A; --gold:#E8C547; --ok:#7CC47C;
}

html, body, .stApp, p, li, label, input, textarea, button, a,
[data-baseweb="select"] *{
  font-family:'Barlow', system-ui, -apple-system, 'Segoe UI', sans-serif;
}
.stApp{ background:var(--bg); color:var(--text); }

.block-container{ max-width:1160px; padding:2.4rem 2rem 4rem; }
@media (max-width:640px){ .block-container{ padding:1.4rem 1rem 3rem; } }

header[data-testid="stHeader"]{ background:transparent; }
#MainMenu, footer, .stDeployButton{ visibility:hidden; }
[data-testid="stHeaderActionElements"]{ display:none; }

/* Marcadores invisíveis (e a linha que eles ocupariam no layout) */
.cm, .hm, .calendar-marker, .recentes-marker, .acoes-marker, .d-marca{ display:none; }
.element-container:has(.hm), .element-container:has(.calendar-marker),
.element-container:has(.recentes-marker), .element-container:has(.acoes-marker),
.element-container:has(.d-marca),
[data-testid="stElementContainer"]:has(.hm),
[data-testid="stElementContainer"]:has(.calendar-marker),
[data-testid="stElementContainer"]:has(.recentes-marker),
[data-testid="stElementContainer"]:has(.acoes-marker),
[data-testid="stElementContainer"]:has(.d-marca){ display:none; }

/* ---------- Tipografia ---------- */
h1, h2, h3, .page-title, .sec, .brand-name{
  font-family:'Barlow Condensed','Barlow',sans-serif !important;
}
.stApp .page-title{
  font-size:3.2rem; font-weight:700; line-height:1; letter-spacing:-.01em;
  margin:0; padding:0; color:var(--text);
}
.stApp .page-sub{ color:var(--muted); font-size:1.05rem; margin:.45rem 0 1.6rem; }
.stApp .sec{
  font-size:1.65rem; font-weight:600; line-height:1.1; margin:2rem 0 .3rem;
  padding:0; color:var(--text);
}
.stApp .sec-sub{ color:var(--muted); font-size:.98rem; margin:0 0 .9rem; }
@media (max-width:640px){ .stApp .page-title{ font-size:2.5rem; } }
[data-testid="stCaptionContainer"]{ color:var(--muted); }

/* ---------- Barra lateral ---------- */
section[data-testid="stSidebar"]{
  background:rgba(0,0,0,.88); border-right:1px solid var(--line);
}
section[data-testid="stSidebar"] > div{ padding-top:.4rem; }
.brand{ padding:.9rem .75rem 1.1rem; }
.brand-bar{ display:block; width:38px; height:5px; border-radius:3px; background:var(--accent); margin-bottom:.7rem; }
.brand-name{ display:block; font-size:1.85rem; font-weight:700; line-height:1; color:var(--text); }
.brand-sub{ display:block; color:var(--muted); font-size:.95rem; margin-top:.3rem; }
.side-title{ color:var(--muted); font-size:.9rem; font-weight:600; padding:1rem .75rem .3rem; }

/* Itens do menu alinhados à esquerda (o texto fica dentro de div > p) */
section[data-testid="stSidebar"] .stButton > button,
section[data-testid="stSidebar"] .stDownloadButton > button{
  justify-content:flex-start; text-align:left; background:transparent;
  border:0; border-left:3px solid transparent; border-radius:6px;
  color:var(--muted); font-weight:500; font-size:1rem;
  padding:.45rem .75rem; box-shadow:none; min-height:2.3rem;
}
section[data-testid="stSidebar"] .stButton > button > div,
section[data-testid="stSidebar"] .stDownloadButton > button > div,
section[data-testid="stSidebar"] .stButton [data-testid="stMarkdownContainer"],
section[data-testid="stSidebar"] .stDownloadButton [data-testid="stMarkdownContainer"]{
  width:100%; justify-content:flex-start; text-align:left;
}
section[data-testid="stSidebar"] .stButton > button p,
section[data-testid="stSidebar"] .stDownloadButton > button p{ text-align:left; width:100%; }
section[data-testid="stSidebar"] .stButton > button:hover,
section[data-testid="stSidebar"] .stDownloadButton > button:hover{
  background:var(--raised); color:var(--text); border-left-color:transparent;
}
section[data-testid="stSidebar"] .stButton > button[kind="primary"],
section[data-testid="stSidebar"] .stButton > button[data-testid="stBaseButton-primary"]{
  background:rgba(225,38,47,.16); color:#fff; border-left:3px solid var(--accent);
}

.sync{ display:flex; align-items:center; gap:.55rem; color:var(--muted); font-size:.92rem; padding:.2rem .75rem .5rem; }
.sync i{ width:9px; height:9px; border-radius:50%; background:var(--line); flex:none; }
.sync.ok i{ background:var(--ok); }
.sync.dirty i{ background:var(--gold); }
.sync.dirty{ color:var(--text); }
.sync-erro{ color:var(--muted); font-size:.85rem; padding:0 .75rem .6rem; line-height:1.35; }

/* ---------- Botões ---------- */
.stApp .stButton > button,
.stApp .stDownloadButton > button,
.stApp [data-testid="stFormSubmitButton"] > button,
.stApp a[data-testid="stBaseLinkButton-secondary"]{
  border-radius:6px; border:1px solid var(--line); background:transparent;
  color:var(--text); font-weight:500; min-height:2.4rem; box-shadow:none;
  transition:background .15s, border-color .15s;
}
.stApp .stButton > button:hover,
.stApp .stDownloadButton > button:hover,
.stApp [data-testid="stFormSubmitButton"] > button:hover,
.stApp a[data-testid="stBaseLinkButton-secondary"]:hover{
  background:var(--raised); border-color:var(--muted); color:var(--text);
}
.stApp .stButton > button[kind="primary"],
.stApp .stButton > button[data-testid="stBaseButton-primary"],
.stApp [data-testid="stFormSubmitButton"] > button[kind="primaryFormSubmit"],
.stApp [data-testid="stFormSubmitButton"] > button[data-testid="stBaseButton-primaryFormSubmit"]{
  background:var(--accent); border-color:var(--accent); color:#fff; font-weight:600;
}
.stApp .stButton > button[kind="primary"]:hover,
.stApp .stButton > button[data-testid="stBaseButton-primary"]:hover,
.stApp [data-testid="stFormSubmitButton"] > button[kind="primaryFormSubmit"]:hover,
.stApp [data-testid="stFormSubmitButton"] > button[data-testid="stBaseButton-primaryFormSubmit"]:hover{
  background:#C42F34; border-color:#C42F34; color:#fff;
}
.stApp button:focus-visible, .stApp a:focus-visible{ outline:2px solid var(--gold); outline-offset:2px; }

/* ---------- Campos, expansores, formulários ---------- */
/* Contorno vermelho estático; o brilho aparece só no foco (sem animação contínua). */
[data-testid="stTextInput"] [data-baseweb="input"],
[data-testid="stTextArea"] [data-baseweb="textarea"],
[data-testid="stNumberInput"] [data-baseweb="input"],
[data-testid="stSelectbox"] [data-baseweb="select"] > div,
[data-testid="stDateInput"] [data-baseweb="input"]{
  border:1px solid rgba(225,38,47,.5) !important; border-radius:6px;
  transition:box-shadow .2s ease, border-color .2s ease;
}
[data-testid="stTextInput"] [data-baseweb="input"]:focus-within,
[data-testid="stTextArea"] [data-baseweb="textarea"]:focus-within,
[data-testid="stNumberInput"] [data-baseweb="input"]:focus-within,
[data-testid="stSelectbox"] [data-baseweb="select"] > div:focus-within,
[data-testid="stDateInput"] [data-baseweb="input"]:focus-within{
  border-color:var(--accent-hi) !important;
  box-shadow:0 0 9px rgba(225,38,47,.35) !important;
}
[data-testid="stExpander"] details,
[data-testid="stForm"]{
  border:1px solid rgba(225,38,47,.45); border-radius:8px;
  background:rgba(8,8,8,.86);
}

/* ---------- Cartões de registro ---------- */
div[data-testid="stVerticalBlockBorderWrapper"]:has(.cm){
  background:rgba(8,8,8,.86); border:1px solid var(--line); border-radius:8px;
}
.card-top{ display:flex; flex-wrap:wrap; align-items:center; gap:.4rem .75rem; margin-bottom:.35rem; }
.card-meta{ color:var(--muted); font-size:.95rem; }
.diff{ display:inline-flex; align-items:center; gap:.4rem; color:var(--muted); font-size:.92rem; }
.pips{ display:inline-flex; gap:3px; }
.pips i{ width:6px; height:12px; border-radius:2px; background:var(--line); }
.pips i.on{ background:var(--text); }
.card-date{ color:var(--muted); font-size:.92rem; margin-left:auto; }
.card-title{
  font-family:'Barlow Condensed','Barlow',sans-serif; font-size:1.7rem;
  font-weight:600; line-height:1.15; color:var(--text); margin:.1rem 0 .1rem;
}
.card-title .fav{ color:var(--gold); margin-right:.3rem; }
.card-src{ color:var(--muted); font-size:.92rem; margin:.2rem 0 .5rem; }

.pill{
  display:inline-block; padding:.1rem .6rem; border-radius:999px;
  font-size:.85rem; font-weight:600; color:#000; background:var(--c, var(--muted));
}

@media (max-width:640px){
  [data-testid="stHorizontalBlock"]:has(.acoes-marker){ flex-wrap:nowrap !important; gap:.35rem !important; }
  [data-testid="stHorizontalBlock"]:has(.acoes-marker) > [data-testid="stColumn"],
  [data-testid="stHorizontalBlock"]:has(.acoes-marker) > [data-testid="column"]{
    min-width:0 !important; flex:1 1 0 !important; width:auto !important;
  }
  [data-testid="stHorizontalBlock"]:has(.acoes-marker) > [data-testid="stColumn"]:last-child,
  [data-testid="stHorizontalBlock"]:has(.acoes-marker) > [data-testid="column"]:last-child{ display:none; }
  [data-testid="stHorizontalBlock"]:has(.acoes-marker) button p{ font-size:.85rem; }
}

/* ---------- Visão geral ---------- */
.pre-item{ padding:.7rem 0; border-top:1px solid rgba(225,38,47,.35); }
.pre-item.first{ border-top:0; }
.pre-title{ font-family:'Barlow Condensed','Barlow',sans-serif; font-size:1.35rem; font-weight:600; line-height:1.15; margin:.35rem 0 .1rem; }
.pre-desc{ color:var(--muted); font-size:.98rem; line-height:1.4; }

.bar-row{ display:grid; grid-template-columns:minmax(110px,170px) 1fr 34px; gap:.7rem; align-items:center; margin:.4rem 0; }
.bar-label{ color:var(--text); font-size:.98rem; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; }
.bar-track{ height:10px; background:var(--raised); border-radius:5px; overflow:hidden; }
.bar-fill{ display:block; height:100%; border-radius:5px; background:var(--c, var(--muted)); }
.bar-n{ text-align:right; color:var(--muted); font-size:.95rem; font-variant-numeric:tabular-nums; }

.recente{ display:flex; flex-wrap:wrap; align-items:center; gap:.3rem .75rem; padding:.65rem .2rem; border-top:1px solid rgba(225,38,47,.35); }
.recente:first-child{ border-top:0; }
.recente .t{ font-weight:500; flex:1 1 220px; min-width:0; }
.recente .d{ color:var(--muted); font-size:.92rem; }

.vazio{ border:1px dashed var(--line); border-radius:8px; padding:1.6rem 1.4rem; color:var(--muted); line-height:1.5; }
.vazio b{ color:var(--text); font-weight:600; }

/* ---------- Indicadores ---------- */
.kpi-card{
  position:relative; box-sizing:border-box; min-height:105px; height:100%; padding:.75rem .85rem;
  background:#000; border:1px solid rgba(225,38,47,.85); border-radius:10px;
}
.kpi-icon{ font-size:1.1rem; color:var(--accent-hi); margin-bottom:.25rem; }
.kpi-value{ font-family:'Barlow Condensed','Barlow',sans-serif; font-size:2rem; font-weight:700; line-height:1.05; color:#fff; }
.kpi-label{ font-size:.88rem; color:var(--muted); margin-top:.25rem; }
@media (max-width:640px){
  [data-testid="stHorizontalBlock"]:has(.kpi-card){ gap:.6rem !important; }
  [data-testid="stHorizontalBlock"]:has(.kpi-card) > [data-testid="stColumn"],
  [data-testid="stHorizontalBlock"]:has(.kpi-card) > [data-testid="column"]{
    min-width:calc(50% - .3rem) !important; flex:1 1 calc(50% - .3rem) !important;
  }
  .kpi-card{ min-height:92px; }
}

/* ---------- Calendário de treinos ---------- */
.cal-month{
  font-family:'Barlow Condensed','Barlow',sans-serif; font-size:1.6rem; font-weight:700;
  text-align:center; line-height:2.4rem; text-transform:capitalize; white-space:nowrap;
}
@media (max-width:640px){ .cal-month{ font-size:1.15rem; } }
.cal-dow{ text-align:center; color:var(--muted); font-size:.78rem; font-weight:700; padding:.25rem 0 .4rem; }
.cal-vazio{ height:2.7rem; }
.cal-legend{ display:flex; gap:1.1rem; flex-wrap:wrap; color:var(--muted); font-size:.9rem; margin:.7rem 0 .15rem; }
.cal-legend i{ display:inline-block; width:12px; height:12px; border-radius:3px; margin-right:.4rem; vertical-align:-1px; }
.cal-legend i.treino{ background:var(--accent); }
.cal-legend i.hoje{ background:transparent; box-shadow:inset 0 0 0 2px #fff; }
.cal-resumo{ color:var(--muted); font-size:.92rem; margin-top:.35rem; }
.cal-resumo b{ color:var(--text); font-weight:600; }
"""


def _fundo_em_base64():
    """Lê boxing_background.png, reduz e comprime (evita mandar vários MB a cada clique)."""
    caminho = Path(__file__).resolve().parent / ARQUIVO_FUNDO

    if not caminho.exists():
        return "", ""

    try:
        from PIL import Image

        with Image.open(caminho) as imagem:
            imagem = imagem.convert("RGB")
            imagem.thumbnail((1920, 1080))
            saida = io.BytesIO()
            imagem.save(saida, "JPEG", quality=72, optimize=True, progressive=True)
        return base64.b64encode(saida.getvalue()).decode("ascii"), "image/jpeg"
    except Exception:
        try:
            return base64.b64encode(caminho.read_bytes()).decode("ascii"), "image/png"
        except Exception:
            return "", ""


@st.cache_data(show_spinner=False)
def montar_css():
    """CSS completo do app. Fica em cache: só é montado uma vez por servidor."""
    # Fita colorida lateral dos cartões e cor das pílulas, por tipo
    por_tipo = ""
    for info in TIPOS_INFO.values():
        classe, cor = info["css"], info["cor"]
        por_tipo += (
            f".{classe}{{--c:{cor};}}"
            f"div[data-testid='stVerticalBlockBorderWrapper']:has(.cm.{classe})"
            f"{{border-left:5px solid {cor} !important;}}"
        )

    # Painéis neon: borda fixa + um brilho que pulsa só em opacidade
    # (barato para o navegador; animar box-shadow redesenha a tela inteira).
    seletores = ",".join(
        f"div[data-testid='stVerticalBlockBorderWrapper']:has(.{m})" for m in PAINEIS_NEON
    )
    brilho = ",".join(
        f"div[data-testid='stVerticalBlockBorderWrapper']:has(.{m})::after" for m in PAINEIS_NEON
    )
    paineis = (
        "@keyframes pulsoNeon{0%,100%{opacity:.35}50%{opacity:1}}"
        f"{seletores}{{position:relative;background:rgba(5,5,5,.86) !important;"
        "border:1px solid rgba(225,38,47,.75) !important;border-radius:10px !important;}"
        f"{brilho}{{content:'';position:absolute;inset:-1px;border-radius:inherit;"
        "pointer-events:none;box-shadow:0 0 12px rgba(225,38,47,.5), inset 0 0 8px rgba(225,38,47,.12);"
        "opacity:.35;animation:pulsoNeon 3.6s ease-in-out infinite;}"
        ".kpi-card::after{content:'';position:absolute;inset:-1px;border-radius:inherit;"
        "pointer-events:none;box-shadow:0 0 10px rgba(225,38,47,.45);"
        "opacity:.35;animation:pulsoNeon 3.6s ease-in-out infinite;}"
        "div[data-testid='stVerticalBlockBorderWrapper']:has(.hm){padding:.35rem .35rem .15rem;}"
        "@media (prefers-reduced-motion: reduce){"
        f"{brilho},.kpi-card::after{{animation:none;opacity:.6;}}}}"
    )

    # Calendário: botões dos dias, hoje, e os 7 dias lado a lado no celular
    hoje_ = (
        ".element-container:has(.d-hoje) + .element-container button,"
        "[data-testid='stElementContainer']:has(.d-hoje) + [data-testid='stElementContainer'] button"
    )
    calendario = (
        f"{CALENDARIO} [data-testid='stHorizontalBlock']{{gap:.35rem;}}"
        f"{CALENDARIO} .stButton button{{min-height:2.7rem;padding:.2rem 0;font-weight:600;"
        "font-variant-numeric:tabular-nums;}"
        f"{CALENDARIO} .stButton button p{{font-size:.98rem;overflow:visible;"
        "text-overflow:clip;white-space:nowrap;}"
        f"{CALENDARIO} .stButton button[kind='primary'],"
        f"{CALENDARIO} .stButton button[data-testid='stBaseButton-primary']"
        "{box-shadow:0 0 8px rgba(225,38,47,.45);}"
        f"{CALENDARIO} .stButton button:disabled{{opacity:.35;}}"
        f"{hoje_}{{box-shadow:inset 0 0 0 2px #fff !important;}}"
        "@media (max-width:640px){"
        f"{CALENDARIO} [data-testid='stHorizontalBlock']{{flex-wrap:nowrap !important;gap:.25rem !important;}}"
        f"{CALENDARIO} [data-testid='stColumn'],{CALENDARIO} [data-testid='column']"
        "{min-width:0 !important;}}"
    )

    # Fundo cinematográfico (opcional)
    fundo = ""
    imagem_b64, mime = _fundo_em_base64()
    if imagem_b64:
        fundo = (
            ".stApp{background-color:#000 !important;"
            "background-image:linear-gradient(rgba(0,0,0,.66),rgba(0,0,0,.72)),"
            f"url('data:{mime};base64,{imagem_b64}') !important;"
            "background-position:center center !important;background-size:cover !important;"
            "background-attachment:fixed !important;background-repeat:no-repeat !important;}"
            "[data-testid='stAppViewContainer'],[data-testid='stMain']{background:transparent !important;}"
            "@media (max-width:640px){.stApp{background-attachment:scroll !important;}}"
        )

    return "<style>" + CSS_BASE + fundo + por_tipo + paineis + calendario + "</style>"


# ============================================================
# BLOCOS DE HTML
# ============================================================

def html_pilula(tipo):
    classe = TIPOS_INFO[tipo]["css"]
    return f'<span class="pill {classe}">{escape(tipo)}</span>'


def html_dificuldade(dificuldade):
    nivel = PESO_DIFICULDADE.get(dificuldade, 2)
    barras = "".join(f'<i class="{"on" if k < nivel else ""}"></i>' for k in range(3))
    return f'<span class="diff"><span class="pips">{barras}</span>{escape(dificuldade)}</span>'


def cabecalho_pagina(titulo, subtitulo=""):
    st.markdown(
        f'<h1 class="page-title">{escape(titulo)}</h1>'
        f'<div class="page-sub">{escape(subtitulo)}</div>',
        unsafe_allow_html=True,
    )


def secao(titulo, subtitulo=""):
    sub = f'<div class="sec-sub">{escape(subtitulo)}</div>' if subtitulo else ""
    st.markdown(f'<h2 class="sec">{escape(titulo)}</h2>{sub}', unsafe_allow_html=True)


def estado_vazio(titulo, texto):
    st.markdown(
        f'<div class="vazio"><b>{escape(titulo)}</b><br>{escape(texto)}</div>',
        unsafe_allow_html=True,
    )


# ============================================================
# BARRA LATERAL
# ============================================================

def contar_para_menu(nome, tipo):
    df = st.session_state.df
    if tipo:
        return int((df["tipo"] == tipo).sum())
    if nome == "Caderno":
        return len(df)
    if nome == "Favoritos":
        return int(df["favorito"].sum())
    return None


def barra_lateral():
    with st.sidebar:
        st.markdown(
            '<div class="brand"><span class="brand-bar"></span>'
            '<span class="brand-name">Caderno tático</span>'
            '<span class="brand-sub">Seu diário de boxe</span></div>',
            unsafe_allow_html=True,
        )

        for nome, tipo in PAGINAS:
            quantidade = contar_para_menu(nome, tipo)
            rotulo = nome if quantidade is None else f"{nome} ({quantidade})"

            if st.button(
                rotulo,
                key=f"nav_{nome}",
                type="primary" if st.session_state.pagina == nome else "secondary",
                **LARGURA_BTN,
            ):
                ir_para(nome)
                st.rerun()

        st.markdown('<div class="side-title">Dados</div>', unsafe_allow_html=True)

        # Situação da sincronização
        if not github_configurado():
            chip = '<div class="sync off"><i></i>Nuvem não configurada</div>'
        elif st.session_state.estado_nuvem == "erro":
            chip = '<div class="sync dirty"><i></i>Sem conexão com a nuvem</div>'
        elif st.session_state.alterado:
            chip = '<div class="sync dirty"><i></i>Alterações não salvas</div>'
        elif st.session_state.ultimo_salvamento:
            chip = f'<div class="sync ok"><i></i>Salvo às {st.session_state.ultimo_salvamento}</div>'
        else:
            chip = '<div class="sync ok"><i></i>Sincronizado com a nuvem</div>'

        st.markdown(chip, unsafe_allow_html=True)

        if st.session_state.erro_conexao and st.session_state.estado_nuvem != "sem_config":
            st.markdown(
                f'<div class="sync-erro">{escape(st.session_state.erro_conexao)}</div>',
                unsafe_allow_html=True,
            )

        if github_configurado() and not _segredo("APP_PASSWORD", ""):
            st.markdown(
                '<div class="sync-erro">Sem senha: qualquer pessoa com o link do app '
                "pode ver e editar o caderno. Defina APP_PASSWORD nos Secrets.</div>",
                unsafe_allow_html=True,
            )

        if st.button("Salvar na nuvem", key="btn_salvar", **LARGURA_BTN):
            salvar_na_nuvem()
            st.rerun()

        # Recarregar pede confirmação se houver alterações pendentes
        if st.button("Recarregar da nuvem", key="btn_recarregar", **LARGURA_BTN):
            if st.session_state.alterado and not st.session_state.confirmar_recarga:
                st.session_state.confirmar_recarga = True
                st.rerun()
            else:
                recarregar_da_nuvem()

        if st.session_state.confirmar_recarga:
            st.warning("Há alterações não salvas. Recarregar vai descartá-las.")
            c1, c2 = st.columns(2)
            if c1.button("Descartar", key="recarga_sim", type="primary", **LARGURA_BTN):
                recarregar_da_nuvem()
            if c2.button("Voltar", key="recarga_nao", **LARGURA_BTN):
                st.session_state.confirmar_recarga = False
                st.rerun()

        if github_configurado():
            st.toggle(
                "Salvar sozinho a cada mudança",
                key="autosave",
                help="Cada alteração gera um commit no repositório do GitHub.",
            )

        with st.expander("Exportar dados"):
            exportar_dados()


def recarregar_da_nuvem():
    df_novo, erro, estado = carregar_github()
    st.session_state.confirmar_recarga = False

    if estado in ("ok", "novo"):
        st.session_state.df = preparar_dataframe(df_novo)
        st.session_state.alterado = False
        st.session_state.estado_nuvem = estado
        st.session_state.erro_conexao = ""
        st.session_state.editando_id = None
        st.session_state.excluindo_id = None
        avisar("Caderno recarregado da nuvem.", "🔄")
    else:
        st.session_state.erro_conexao = erro
        avisar(erro or "Não foi possível recarregar.", "⚠️")

    st.rerun()


def exportar_dados():
    st.caption("Exporte seus registros para guardar uma cópia de segurança.")
    try:
        st.download_button(
            "Baixar planilha (.xlsx)",
            data=gerar_excel(st.session_state.df),
            file_name=f"caderno_tatico_{hoje():%Y-%m-%d}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="baixar_excel",
            **LARGURA_BTN,
        )
        st.download_button(
            "Baixar backup (.json)",
            data=gerar_backup_json(st.session_state.df),
            file_name=f"caderno_tatico_{hoje():%Y-%m-%d}.json",
            mime="application/json",
            key="baixar_json",
            **LARGURA_BTN,
        )
    except Exception as erro:
        st.error(f"Erro ao gerar os arquivos: {erro}")


# ============================================================
# FORMULÁRIOS
# ============================================================

def campos_do_registro(prefixo, base):
    """Desenha os campos dentro de um st.form e devolve os valores digitados."""
    c1, c2 = st.columns(2)

    with c1:
        tipo = st.selectbox(
            "Tipo", TIPOS, index=indice_de(TIPOS, base["tipo"]), key=f"{prefixo}_tipo"
        )
        titulo = st.text_input(
            "Título *",
            value=base["titulo"],
            placeholder="Ex.: Jab duplo seguido de direto",
            key=f"{prefixo}_titulo",
        )
        categorias = CATEGORIAS + (
            [base["categoria"]] if base["categoria"] not in CATEGORIAS else []
        )
        categoria = st.selectbox(
            "Categoria",
            categorias,
            index=indice_de(categorias, base["categoria"], len(categorias) - 1),
            key=f"{prefixo}_categoria",
        )
        dificuldade = st.selectbox(
            "Dificuldade",
            DIFICULDADES,
            index=indice_de(DIFICULDADES, base["dificuldade"], 1),
            key=f"{prefixo}_dificuldade",
        )

    with c2:
        data_registro = st.date_input(
            "Data",
            value=para_date(base["data"]) if base["data"] else hoje(),
            format="DD/MM/YYYY",
            key=f"{prefixo}_data",
        )
        fonte = st.text_input(
            "Fonte",
            value=base["fonte"],
            placeholder="Professor, vídeo, sparring com...",
            key=f"{prefixo}_fonte",
        )
        link_video = st.text_input(
            "Link do vídeo",
            value=base["link_video"],
            placeholder="https://",
            key=f"{prefixo}_link",
        )

    favorito = st.checkbox(
        "Marcar como favorito", value=bool(base["favorito"]), key=f"{prefixo}_favorito"
    )
    descricao = st.text_area(
        "Descrição *",
        value=base["descricao"],
        height=170,
        placeholder="O passo a passo, o que deu errado e como corrigir. Aceita Markdown, como listas com '-'.",
        key=f"{prefixo}_descricao",
    )

    return {
        "tipo": tipo,
        "titulo": titulo.strip(),
        "categoria": categoria,
        "descricao": descricao.strip(),
        "fonte": fonte.strip(),
        "link_video": normalizar_link(link_video),
        "dificuldade": dificuldade,
        "favorito": favorito,
        "data": data_registro.strftime("%Y-%m-%d"),
    }


def validar_campos(campos):
    if not campos["titulo"]:
        return "Dê um título ao registro."
    if not campos["descricao"]:
        return "Escreva a descrição do registro."
    return ""


def formulario_novo(chave, tipo_padrao=None, expandido=False, data_padrao="", rotulo="Novo registro"):
    base = {**VALORES_PADRAO, "tipo": tipo_padrao or "Dica", "data": data_padrao}

    with st.expander(rotulo, expanded=expandido):
        with st.form(f"form_novo_{chave}", clear_on_submit=True):
            campos = campos_do_registro(f"novo_{chave}", base)
            enviar = st.form_submit_button("Adicionar ao caderno", type="primary", **LARGURA_BTN)

        if enviar:
            erro = validar_campos(campos)
            if erro:
                st.error(erro)
            else:
                adicionar_registro(campos)
                st.rerun()


def formulario_edicao(registro, chave):
    with st.form(f"form_edicao_{chave}_{registro['id']}"):
        campos = campos_do_registro(f"ed_{chave}_{registro['id']}", registro)

        c1, c2 = st.columns(2)
        salvar = c1.form_submit_button("Salvar alterações", type="primary", **LARGURA_BTN)
        cancelar = c2.form_submit_button("Cancelar", **LARGURA_BTN)

    if cancelar:
        st.session_state.editando_id = None
        st.rerun()

    if salvar:
        erro = validar_campos(campos)
        if erro:
            st.error(erro)
        else:
            st.session_state.editando_id = None
            atualizar_registro(registro["id"], campos, "Registro atualizado.")
            st.rerun()


# ============================================================
# CARTÃO DE REGISTRO
# ============================================================

def cartao_registro(registro, chave):
    id_registro = registro["id"]
    tipo = registro["tipo"]
    info = TIPOS_INFO[tipo]

    with st.container(border=True):
        if st.session_state.editando_id == id_registro:
            st.markdown(
                f'<span class="cm {info["css"]}"></span>'
                f'<div class="card-title">Editando: {escape(registro["titulo"])}</div>',
                unsafe_allow_html=True,
            )
            formulario_edicao(registro, chave)
            return

        estrela = '<span class="fav">★</span>' if registro["favorito"] else ""
        data_html = (
            f'<span class="card-date">{formatar_data(registro["data"])}</span>'
            if registro["data"]
            else ""
        )

        st.markdown(
            f'<span class="cm {info["css"]}"></span>'
            f'<div class="card-top">{html_pilula(tipo)}'
            f'<span class="card-meta">{escape(registro["categoria"])}</span>'
            f'{html_dificuldade(registro["dificuldade"])}{data_html}</div>'
            f'<div class="card-title">{estrela}{escape(registro["titulo"])}</div>',
            unsafe_allow_html=True,
        )

        if registro["descricao"]:
            st.markdown(registro["descricao"])

        if registro["fonte"]:
            st.markdown(
                f'<div class="card-src">Fonte: {escape(registro["fonte"])}</div>',
                unsafe_allow_html=True,
            )

        # Confirmação de exclusão
        if st.session_state.excluindo_id == id_registro:
            st.warning("Excluir este registro? Essa ação não pode ser desfeita.")
            c1, c2, _ = st.columns([1, 1, 2])
            if c1.button("Excluir", key=f"{chave}_conf_{id_registro}", type="primary", **LARGURA_BTN):
                st.session_state.excluindo_id = None
                excluir_registro(id_registro)
                st.rerun()
            if c2.button("Cancelar", key=f"{chave}_canc_{id_registro}", **LARGURA_BTN):
                st.session_state.excluindo_id = None
                st.rerun()
            return

        acoes_do_cartao(registro, chave)


def acoes_do_cartao(registro, chave):
    id_registro = registro["id"]
    tem_video = bool(registro["link_video"])

    quantidade = 3 + int(tem_video)
    colunas = st.columns([1] * quantidade + [max(1, 6 - quantidade)])
    posicao = 0

    # Marcador invisível: permite manter os botões numa linha só no celular
    with colunas[0]:
        st.markdown('<span class="acoes-marker"></span>', unsafe_allow_html=True)

    if tem_video:
        with colunas[posicao]:
            st.link_button("Vídeo", registro["link_video"], **LARGURA_BTN)
        posicao += 1

    with colunas[posicao]:
        if st.button(
            "★" if registro["favorito"] else "☆",
            key=f"{chave}_fav_{id_registro}",
            help="Tirar dos favoritos" if registro["favorito"] else "Adicionar aos favoritos",
            **LARGURA_BTN,
        ):
            atualizar_registro(
                id_registro,
                {"favorito": not registro["favorito"]},
                "Removido dos favoritos." if registro["favorito"] else "Adicionado aos favoritos.",
                "⭐",
            )
            st.rerun()
    posicao += 1

    with colunas[posicao]:
        if st.button("Editar", key=f"{chave}_edit_{id_registro}", **LARGURA_BTN):
            st.session_state.editando_id = id_registro
            st.session_state.excluindo_id = None
            st.rerun()
    posicao += 1

    with colunas[posicao]:
        if st.button("Excluir", key=f"{chave}_del_{id_registro}", **LARGURA_BTN):
            st.session_state.excluindo_id = id_registro
            st.session_state.editando_id = None
            st.rerun()


# ============================================================
# LISTAGEM COM FILTROS
# ============================================================

def aplicar_filtros(df, busca, categoria, dificuldade, so_favoritos, tipo, ordem):
    resultado = df.assign(_pos=df.index)

    if tipo != "Todos":
        resultado = resultado[resultado["tipo"] == tipo]
    if categoria != "Todas":
        resultado = resultado[resultado["categoria"] == categoria]
    if dificuldade != "Todas":
        resultado = resultado[resultado["dificuldade"] == dificuldade]
    if so_favoritos:
        resultado = resultado[resultado["favorito"]]

    termos = normalizar_texto(busca).split()
    if termos and len(resultado):
        palheiro = (
            resultado["titulo"] + " " + resultado["descricao"] + " "
            + resultado["categoria"] + " " + resultado["fonte"] + " "
            + resultado["tipo"]
        ).map(normalizar_texto)
        mascara = pd.Series(True, index=resultado.index)
        for termo in termos:
            mascara &= palheiro.str.contains(termo, regex=False)
        resultado = resultado[mascara]

    if ordem == "Mais antigos":
        resultado = resultado.sort_values(["data", "_pos"])
    elif ordem == "Título (A a Z)":
        resultado = resultado.assign(_t=resultado["titulo"].map(normalizar_texto)).sort_values("_t")
    elif ordem == "Mais difíceis primeiro":
        resultado = resultado.assign(_d=resultado["dificuldade"].map(PESO_DIFICULDADE)).sort_values(
            ["_d", "_pos"], ascending=[False, False]
        )
    elif ordem == "Favoritos primeiro":
        resultado = resultado.sort_values(["favorito", "_pos"], ascending=[False, False])
    else:
        resultado = resultado.sort_values(["data", "_pos"], ascending=[False, False])

    return resultado


def lista_de_registros(df_base, chave, filtrar_tipo=False, vazio=None):
    """Mostra filtros e cartões para o DataFrame recebido."""
    if len(df_base) == 0:
        titulo, texto = vazio or ("Nada por aqui ainda.", "Adicione um registro no botão acima.")
        estado_vazio(titulo, texto)
        return

    busca = st.text_input(
        "Buscar",
        placeholder="Buscar por título, descrição, fonte ou categoria",
        label_visibility="collapsed",
        key=f"busca_{chave}",
    )

    posicoes = iter(st.columns([2, 2, 2, 2, 2] if filtrar_tipo else [2, 2, 2, 2]))

    tipo = "Todos"
    if filtrar_tipo:
        with next(posicoes):
            tipo = st.selectbox("Tipo", ["Todos"] + TIPOS, key=f"tipo_{chave}")
    with next(posicoes):
        categorias = ["Todas"] + sorted(c for c in df_base["categoria"].unique() if c)
        categoria = st.selectbox("Categoria", categorias, key=f"cat_{chave}")
    with next(posicoes):
        dificuldade = st.selectbox("Dificuldade", ["Todas"] + DIFICULDADES, key=f"dif_{chave}")
    with next(posicoes):
        ordem = st.selectbox("Ordenar por", ORDENACOES, key=f"ord_{chave}")

    so_favoritos = st.toggle("Só favoritos", key=f"fav_{chave}")

    resultado = aplicar_filtros(df_base, busca, categoria, dificuldade, so_favoritos, tipo, ordem)

    st.caption(plural(len(resultado), "registro encontrado", "registros encontrados"))

    if len(resultado) == 0:
        estado_vazio("Nenhum registro com esses filtros.", "Limpe a busca ou escolha outras opções.")
        return

    chave_limite = f"limite_{chave}"
    limite = st.session_state.get(chave_limite, ITENS_POR_PAGINA)

    for _, registro in resultado.head(limite).iterrows():
        cartao_registro(registro, chave)

    if len(resultado) > limite:
        restantes = len(resultado) - limite
        if st.button(f"Mostrar mais ({restantes})", key=f"mais_{chave}", **LARGURA_BTN):
            st.session_state[chave_limite] = limite + ITENS_POR_PAGINA
            st.rerun()


# ============================================================
# VISÃO GERAL
# ============================================================

def pagina_visao_geral():
    df = st.session_state.df
    data_extenso = data_por_extenso(hoje())

    if len(df) == 0:
        cabecalho_pagina("Caderno tático de boxe", f"{data_extenso}.")
        estado_vazio(
            "Seu caderno está vazio.",
            "Registre a primeira dica, combo ou correção do professor. "
            "Depois do primeiro sparring, anote os erros: eles aparecem aqui "
            "para você rever antes do próximo. Para começar a contar seus treinos, "
            "toque num dia do calendário.",
        )
        formulario_novo("visao", expandido=True)
    else:
        cabecalho_pagina(
            "Caderno tático de boxe",
            f"{data_extenso}. {plural(len(df), 'registro', 'registros')}, "
            f"{plural(int(df['favorito'].sum()), 'favorito', 'favoritos')}.",
        )
        formulario_novo("visao")
        bloco_indicadores(df)

    esquerda, direita = st.columns([3, 2], gap="large")

    with esquerda:
        bloco_antes_do_sparring(df)
        bloco_recentes(df)

    with direita:
        bloco_calendario_treinos(df)
        if len(df):
            bloco_por_tipo(df)
            bloco_por_categoria(df)


def bloco_indicadores(df):
    """Cards-resumo rápidos."""
    mes_atual = hoje().strftime("%Y-%m")
    dias_no_mes = sum(1 for d in datas_de_treino(df) if d.startswith(mes_atual))

    cards = [
        ("🥊", "Dias treinados neste mês", dias_no_mes),
        ("⚡", "Combos anotados", int((df["tipo"] == "Combo").sum())),
        ("🎯", "Revisões de sparring", int(df["tipo"].isin(TIPOS_REVISAO).sum())),
        ("★", "Favoritos", int(df["favorito"].sum())),
    ]

    for coluna, (icone, rotulo, valor) in zip(st.columns(4, gap="small"), cards):
        with coluna:
            st.markdown(
                f'<div class="kpi-card"><div class="kpi-icon">{icone}</div>'
                f'<div class="kpi-value">{valor}</div>'
                f'<div class="kpi-label">{escape(rotulo)}</div></div>',
                unsafe_allow_html=True,
            )


def bloco_antes_do_sparring(df):
    pendentes = df[df["tipo"].isin(TIPOS_REVISAO)]
    pendentes = pendentes.assign(_pos=pendentes.index).sort_values(
        ["data", "_pos"], ascending=[False, False]
    )

    secao(
        "Antes do sparring",
        "Erros e correções para revisar antes do próximo sparring, do mais recente ao mais antigo.",
    )

    with st.container(border=True):
        st.markdown('<span class="hm"></span>', unsafe_allow_html=True)

        if len(pendentes) == 0:
            st.markdown(
                '<div class="pre-desc" style="padding:.6rem .4rem">'
                "Nada para revisar. Depois do próximo sparring, registre o que deu errado "
                "como <b>Erro no Sparring</b> e ele aparece aqui.</div>",
                unsafe_allow_html=True,
            )
            return

        for numero, (_, item) in enumerate(pendentes.head(4).iterrows()):
            st.markdown(
                f'<div class="pre-item{" first" if numero == 0 else ""}">{html_pilula(item["tipo"])}'
                f'<div class="pre-title">{escape(item["titulo"])}</div>'
                f'<div class="pre-desc">{escape(trecho(item["descricao"]))}</div></div>',
                unsafe_allow_html=True,
            )

        if len(pendentes) > 4:
            st.caption(f"Mais {len(pendentes) - 4} nas páginas de erros e correções.")


def bloco_recentes(df):
    if len(df) == 0:
        return

    secao("Últimos registros")

    recentes = df.assign(_pos=df.index).sort_values(
        ["data", "_pos"], ascending=[False, False]
    ).head(6)

    linhas = ""
    for _, r in recentes.iterrows():
        estrela = "★ " if r["favorito"] else ""
        linhas += (
            f'<div class="recente">{html_pilula(r["tipo"])}'
            f'<span class="t">{estrela}{escape(r["titulo"])}</span>'
            f'<span class="d">{formatar_data(r["data"])}</span></div>'
        )

    with st.container(border=True):
        st.markdown('<span class="recentes-marker"></span>', unsafe_allow_html=True)
        st.markdown(linhas, unsafe_allow_html=True)


def bloco_por_tipo(df):
    secao("Por tipo")

    contagem = df["tipo"].value_counts()
    maximo = max(int(contagem.max()), 1) if len(contagem) else 1

    linhas = ""
    for tipo, info in TIPOS_INFO.items():
        n = int(contagem.get(tipo, 0))
        largura = round(100 * n / maximo)
        linhas += (
            f'<div class="bar-row"><span class="bar-label">{escape(tipo)}</span>'
            f'<span class="bar-track"><span class="bar-fill {info["css"]}" style="width:{largura}%"></span></span>'
            f'<span class="bar-n">{n}</span></div>'
        )

    st.markdown(linhas, unsafe_allow_html=True)


def bloco_por_categoria(df):
    secao("Categorias mais anotadas")

    contagem = df["categoria"].value_counts().head(5)
    maximo = max(int(contagem.max()), 1) if len(contagem) else 1

    linhas = ""
    for categoria, n in contagem.items():
        largura = round(100 * n / maximo)
        linhas += (
            f'<div class="bar-row"><span class="bar-label">{escape(categoria)}</span>'
            f'<span class="bar-track"><span class="bar-fill" style="width:{largura}%;--c:#98A9BB"></span></span>'
            f'<span class="bar-n">{int(n)}</span></div>'
        )

    st.markdown(linhas, unsafe_allow_html=True)


# ============================================================
# CALENDÁRIO DE TREINOS
# ============================================================

def mudar_mes(mes, deslocamento):
    """Primeiro dia do mês anterior (-1) ou seguinte (+1)."""
    if deslocamento < 0:
        return (mes - timedelta(days=1)).replace(day=1)
    return (mes.replace(day=28) + timedelta(days=4)).replace(day=1)


def bloco_calendario_treinos(df):
    """Mês em grade: dias treinados em vermelho; tocar num dia abre o treino daquela data."""
    secao("Calendário de treinos", "Toque num dia para abrir ou registrar o treino daquela data.")

    mes = st.session_state.mes_calendario
    hoje_ = hoje()
    mes_de_hoje = hoje_.replace(day=1)
    datas_treino = datas_de_treino(df)

    with st.container(border=True):
        st.markdown('<span class="calendar-marker"></span>', unsafe_allow_html=True)

        # Navegação entre meses (não passa do mês atual)
        c_ant, c_titulo, c_prox = st.columns([1, 4, 1])

        with c_ant:
            if st.button("‹", key="cal_mes_anterior", help="Mês anterior", **LARGURA_BTN):
                st.session_state.mes_calendario = mudar_mes(mes, -1)
                st.rerun()
        with c_titulo:
            st.markdown(
                f'<div class="cal-month">{MESES[mes.month - 1]} {mes.year}</div>',
                unsafe_allow_html=True,
            )
        with c_prox:
            if st.button(
                "›", key="cal_proximo_mes", help="Próximo mês",
                disabled=mes >= mes_de_hoje, **LARGURA_BTN,
            ):
                st.session_state.mes_calendario = mudar_mes(mes, +1)
                st.rerun()

        # Cabeçalho dos dias da semana (semana começa na segunda)
        for coluna, nome in zip(st.columns(7, gap="small"), ["SEG", "TER", "QUA", "QUI", "SEX", "SÁB", "DOM"]):
            with coluna:
                st.markdown(f'<div class="cal-dow">{nome}</div>', unsafe_allow_html=True)

        for semana in calendar.Calendar(firstweekday=0).monthdayscalendar(mes.year, mes.month):
            for coluna, dia in zip(st.columns(7, gap="small"), semana):
                with coluna:
                    if dia == 0:
                        st.markdown('<div class="cal-vazio"></div>', unsafe_allow_html=True)
                        continue

                    data_dia = date(mes.year, mes.month, dia)
                    iso = data_dia.isoformat()
                    tem_treino = iso in datas_treino

                    if data_dia == hoje_:
                        st.markdown('<span class="d-marca d-hoje"></span>', unsafe_allow_html=True)

                    if st.button(
                        str(dia),
                        key=f"cal_dia_{iso}",
                        help=f"{data_dia:%d/%m/%Y}: " + (
                            "abrir treino registrado" if tem_treino else "registrar treino neste dia"
                        ),
                        type="primary" if tem_treino else "secondary",
                        disabled=data_dia > hoje_,
                        **LARGURA_BTN,
                    ):
                        ir_para("Treinos", data_treino=iso)
                        st.rerun()

        dias_no_mes = sum(1 for d in datas_treino if d.startswith(f"{mes.year}-{mes.month:02d}"))
        sequencia = sequencia_de_treinos(datas_treino, hoje_)

        junto = lambda texto: texto.replace(" ", "\u00a0")  # não quebra "2 dias" no meio
        resumo = (
            f"<b>{junto(plural(dias_no_mes, 'dia treinado', 'dias treinados'))}</b> em {MESES[mes.month - 1]}"
        )
        if sequencia and mes == mes_de_hoje:
            resumo += f". Sequência atual: <b>{junto(plural(sequencia, 'dia', 'dias'))}</b>"

        st.markdown(
            '<div class="cal-legend"><span><i class="treino"></i>Dia treinado</span>'
            '<span><i class="hoje"></i>Hoje</span></div>'
            f'<div class="cal-resumo">{resumo}.</div>',
            unsafe_allow_html=True,
        )

        if mes != mes_de_hoje:
            if st.button("Voltar para o mês atual", key="cal_mes_atual", **LARGURA_BTN):
                st.session_state.mes_calendario = mes_de_hoje
                st.rerun()


# ============================================================
# PÁGINAS
# ============================================================

def subtitulo_dados(df):
    return (
        f"{plural(len(df), 'registro', 'registros')}, "
        f"{plural(int(df['favorito'].sum()), 'favorito', 'favoritos')}"
    )


def pagina_caderno():
    df = st.session_state.df
    cabecalho_pagina("Caderno", subtitulo_dados(df))
    formulario_novo("caderno", expandido=len(df) == 0)
    lista_de_registros(
        df,
        "caderno",
        filtrar_tipo=True,
        vazio=("Seu caderno está vazio.", "Use o botão 'Novo registro' para anotar a primeira dica."),
    )


MENSAGENS_VAZIAS = {
    "Combo": "Anote as sequências que você quer automatizar, como 'jab, direto, gancho de esquerda'.",
    "Defesa": "Registre esquivas, bloqueios e contra-ataques que funcionam para você.",
    "Erro no Sparring": "Depois de cada sparring, anote o que deu errado. Os erros aparecem na Visão geral, em 'Antes do sparring'.",
    "Correção do Professor": "Anote cada correção que o professor fizer, enquanto ainda estiver fresca na memória.",
    "Treino": "Registre seus treinos para acompanhar a evolução.",
    "Dica": "Guarde dicas de técnica, estratégia e mentalidade.",
}


def pagina_treino_do_dia(df, iso):
    """Treinos de uma data escolhida no calendário, com marcação rápida."""
    dia = para_date(iso)
    chave = f"treino_{iso}"
    df_dia = df[(df["tipo"] == "Treino") & (df["data"] == iso)]

    cabecalho_pagina(
        f"Treino de {formatar_data(iso)}",
        f"{data_por_extenso(dia)}. "
        + (
            plural(len(df_dia), "treino registrado", "treinos registrados")
            if len(df_dia)
            else "Nenhum treino registrado neste dia."
        ),
    )

    if st.button("← Voltar para todos os treinos", key="voltar_todos_treinos"):
        ir_para("Treinos")
        st.rerun()

    if len(df_dia) == 0:
        estado_vazio(
            "Treinou neste dia?",
            "Marque o dia como treinado agora e complete os detalhes depois, ou registre o treino completo no formulário.",
        )
        if st.button("Marcar como treinado", key=f"marcar_{iso}", type="primary"):
            marcar_dia_como_treinado(iso)
            st.rerun()

    formulario_novo(
        chave,
        tipo_padrao="Treino",
        expandido=False,
        data_padrao=iso,
        rotulo="Registrar treino com detalhes" if len(df_dia) == 0 else "Adicionar outro treino neste dia",
    )

    for _, registro in df_dia.iterrows():
        cartao_registro(registro, chave)


def pagina_tipo(nome, tipo):
    df = st.session_state.df
    data_selecionada = st.session_state.data_treino_selecionada if tipo == "Treino" else ""

    if data_selecionada:
        pagina_treino_do_dia(df, data_selecionada)
        return

    df_tipo = df[df["tipo"] == tipo]

    cabecalho_pagina(nome, subtitulo_dados(df_tipo))
    formulario_novo(f"tipo_{tipo}", tipo_padrao=tipo, expandido=len(df_tipo) == 0)

    lista_de_registros(
        df_tipo,
        f"tipo_{tipo}",
        vazio=(f"Nenhum registro em {nome.lower()} ainda.", MENSAGENS_VAZIAS.get(tipo, "")),
    )


def pagina_favoritos():
    df = st.session_state.df
    favoritos = df[df["favorito"]]

    cabecalho_pagina("Favoritos", plural(len(favoritos), "registro marcado", "registros marcados"))
    lista_de_registros(
        favoritos,
        "favoritos",
        filtrar_tipo=True,
        vazio=(
            "Nenhum favorito ainda.",
            "Toque na estrela (☆) de qualquer registro para guardá-lo aqui.",
        ),
    )


# ============================================================
# EXECUÇÃO
# ============================================================

st.markdown(montar_css(), unsafe_allow_html=True)
exigir_senha()
iniciar_estado()
barra_lateral()
mostrar_avisos()

_pagina = st.session_state.pagina
_tipo_da_pagina = dict(PAGINAS).get(_pagina)

if _pagina == "Caderno":
    pagina_caderno()
elif _pagina == "Favoritos":
    pagina_favoritos()
elif _tipo_da_pagina:
    pagina_tipo(_pagina, _tipo_da_pagina)
else:
    pagina_visao_geral()
